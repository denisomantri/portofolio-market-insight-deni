import unittest
from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from scraper_market_intelligence import (
    analyze_captions,
    build_report_bytes,
    clean_captions,
    clean_captions_with_metrics,
    clean_records_with_metrics,
    compare_segments,
    extract_captions,
    extract_captions_from_sheets,
    find_segment_column,
    frame_to_records,
    limit_records_per_segment,
    make_experiments,
    wilson_interval,
)


class MarketIntelligenceTests(unittest.TestCase):
    def test_cleaning_keeps_relevant_unique_captions(self):
        captions = [
            "Lakban bening kuat untuk membungkus paket dengan rapi setiap hari",
            "Lakban bening kuat untuk membungkus paket dengan rapi setiap hari",
            "Tas wanita aesthetic cocok untuk acara santai bersama keluarga",
            "Deni started following you and sent a notification today",
        ]

        self.assertEqual(
            clean_captions(captions, "lakban"),
            ["Lakban bening kuat untuk membungkus paket dengan rapi setiap hari"],
        )

    def test_keyword_frequency_counts_unique_captions_not_repeated_words(self):
        captions = [
            "Lakban bening untuk paket, lakban bening kuat dan mudah dipakai",
            "Lakban bening membantu membungkus paket dengan rapi setiap hari",
        ]

        _, keywords, _ = analyze_captions(captions, "lakban")
        keyword_counts = {row["Kata Kunci"]: row["Jumlah Caption yang Menyebut"] for row in keywords}

        self.assertEqual(keyword_counts["bening"], 2)
        self.assertNotIn("lakban", keyword_counts)

    def test_quality_metrics_explain_relevance_and_duplicates(self):
        caption = "Lakban bening kuat untuk membungkus paket dengan rapi setiap hari"
        cleaned, metrics = clean_captions_with_metrics(
            [caption, caption, "Tas wanita aesthetic cocok untuk acara santai bersama keluarga"],
            "lakban",
        )

        self.assertEqual(len(cleaned), 1)
        self.assertEqual(metrics["Kandidat dari sumber"], 3)
        self.assertEqual(metrics["Relevan dengan kata kunci"], 2)
        self.assertEqual(metrics["Duplikat yang dihapus"], 1)

    def test_same_caption_is_preserved_for_different_competitors(self):
        caption = "Lakban bening kuat untuk membungkus paket dengan rapi setiap hari"
        records, metrics = clean_records_with_metrics(
            [
                {"caption": caption, "segment": "Toko A"},
                {"caption": caption, "segment": "Toko B"},
            ],
            "lakban",
            "segment",
        )

        self.assertEqual(len(records), 2)
        self.assertEqual(metrics["Duplikat yang dihapus"], 0)

    def test_missing_competitor_labels_are_reported_and_excluded(self):
        records, metrics = clean_records_with_metrics(
            [
                {"caption": "Lakban bening kuat untuk membungkus paket dengan rapi", "segment": "Toko A"},
                {"caption": "Lakban murah kuat untuk membungkus paket dengan rapi", "segment": ""},
            ],
            "lakban",
            "segment",
        )

        self.assertEqual(len(records), 1)
        self.assertEqual(metrics["Baris tanpa label kompetitor"], 1)

    def test_sample_limit_keeps_competitor_groups_balanced(self):
        records = [
            {"caption": f"Caption lakban kuat dan berguna untuk paket nomor {index}", "segment": "Toko A"}
            for index in range(4)
        ] + [
            {"caption": f"Caption lakban murah dan bagus untuk paket nomor {index}", "segment": "Toko B"}
            for index in range(4)
        ]

        selected = limit_records_per_segment(records, limit=4)

        self.assertEqual([record["segment"] for record in selected], [
            "Toko A", "Toko B", "Toko A", "Toko B",
        ])

    def test_competitor_comparison_reports_caption_share_by_theme(self):
        records = [
            {"caption": "Lakban murah promo untuk membungkus paket dengan rapi", "segment": "Toko A"},
            {"caption": "Lakban promo harga murah untuk kebutuhan kirim paket", "segment": "Toko A"},
            {"caption": "Lakban bening kuat untuk merekatkan kabel di rumah", "segment": "Toko B"},
        ]

        comparison = compare_segments(records)
        a_promo = next(
            row for row in comparison
            if row["Kompetitor / Segmen"] == "Toko A" and row["Tema"] == "Harga dan promo"
        )
        b_promo = next(
            row for row in comparison
            if row["Kompetitor / Segmen"] == "Toko B" and row["Tema"] == "Harga dan promo"
        )

        self.assertEqual(a_promo["Porsi Caption (%)"], 100.0)
        self.assertEqual(b_promo["Porsi Caption (%)"], 0.0)
        self.assertLessEqual(a_promo["Selang Wilson 95% bawah (%)"], 100.0)
        self.assertGreaterEqual(a_promo["Selang Wilson 95% atas (%)"], 100.0)
        self.assertEqual(
            b_promo["Selisih vs Toko A (pp)"],
            -100.0,
        )

    def test_wilson_interval_is_bounded_and_wider_for_small_samples(self):
        small_interval = wilson_interval(5, 10)
        large_interval = wilson_interval(500, 1000)

        self.assertLessEqual(small_interval[0], 0.5)
        self.assertGreaterEqual(small_interval[1], 0.5)
        self.assertGreater(small_interval[1] - small_interval[0], large_interval[1] - large_interval[0])
        self.assertEqual(wilson_interval(0, 0), (0.0, 0.0))

    def test_small_sample_does_not_generate_content_test_claim(self):
        experiments = make_experiments([], sample_count=3)

        self.assertEqual(experiments[0]["Prioritas"], "Kumpulkan data tambahan")
        self.assertIn("Belum cukup bukti", experiments[0]["Status"])

    def test_two_observed_themes_produce_a_grounded_experiment(self):
        themes = [
            {"Tema": "Harga dan promo", "Jumlah Caption": 6},
            {"Tema": "Fungsi dan penggunaan", "Jumlah Caption": 4},
        ]

        experiment = make_experiments(themes, sample_count=20)[0]

        self.assertIn("6 caption", experiment["Dasar temuan"])
        self.assertIn("4 caption", experiment["Dasar temuan"])
        self.assertIn("Hipotesis", experiment["Status"])

    def test_extract_captions_accepts_existing_workbook_column_names(self):
        captions = extract_captions(pd.DataFrame({
            "Sampel Caption / Teks Pasar": ["Lakban bening kuat untuk paket"],
        }))

        self.assertEqual(captions, ["Lakban bening kuat untuk paket"])

    def test_excel_report_upload_skips_summary_sheet_for_caption_sheet(self):
        captions, sheet_name = extract_captions_from_sheets({
            "Ringkasan Keputusan": pd.DataFrame({"Indikator": ["Topik"], "Nilai": ["Lakban"]}),
            "Sampel Terklasifikasi": pd.DataFrame({
                "Sampel Caption": ["Lakban bening kuat untuk paket"],
            }),
        })

        self.assertEqual(sheet_name, "Sampel Terklasifikasi")
        self.assertEqual(captions, ["Lakban bening kuat untuk paket"])

    def test_frame_to_records_preserves_competitor_segment(self):
        frame = pd.DataFrame({
            "Sampel Caption": ["Lakban bening kuat untuk paket"],
            "Kompetitor": ["Toko A"],
        })
        segment_column = find_segment_column(frame)
        records, caption_column = frame_to_records(frame, segment_column)

        self.assertEqual(caption_column, "Sampel Caption")
        self.assertEqual(records[0]["segment"], "Toko A")

    def test_report_bytes_contains_decision_and_analysis_sheets(self):
        captions = [
            "Lakban murah untuk membungkus paket dengan sangat rapi setiap hari",
            "Lakban bening dan kuat untuk merekatkan kabel dengan mudah sekali",
        ]
        _, metrics = clean_captions_with_metrics(captions, "lakban")

        workbook = load_workbook(BytesIO(build_report_bytes(
            "lakban",
            "Unit test",
            50,
            captions,
            metrics,
            comparison_rows=[{
                "Kompetitor / Segmen": "Toko A",
                "Tema": "Harga dan promo",
                "Porsi Caption (%)": 100.0,
            }],
        )))

        self.assertIn("Ringkasan Keputusan", workbook.sheetnames)
        self.assertIn("Eksperimen Berikutnya", workbook.sheetnames)
        self.assertIn("Perbandingan Kompetitor", workbook.sheetnames)


if __name__ == "__main__":
    unittest.main()
